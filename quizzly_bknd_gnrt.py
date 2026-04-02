import os
import PyPDF2
from pptx import Presentation
import docx
from openai import OpenAI
from langchain_openai import ChatOpenAI
from langchain_core.messages import SystemMessage, HumanMessage
from langchain_core.output_parsers import JsonOutputParser

def setup_api():
    """Validates the API key and returns the native OpenAI client."""
    api_key = os.environ.get('OPENAI_API_KEY')
    if not api_key:
        raise ValueError("OPENAI_API_KEY not found. Please set it in your environment.")
    return OpenAI(api_key=api_key)

def get_page_count(file_path):
    """Calculates the number of pages/slides in the uploaded document."""
    ext = os.path.splitext(file_path)[1].lower()
    
    try:
        if ext == '.pdf':
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                return len(reader.pages)
                
        elif ext == '.pptx':
            prs = Presentation(file_path)
            return len(prs.slides)
            
        elif ext == '.docx':
            # Word docs don't have strict pages, so we estimate ~500 words per page
            doc = docx.Document(file_path)
            word_count = sum(len(p.text.split()) for p in doc.paragraphs)
            return max(1, word_count // 500)
            
        else:
            return 1 # Fallback
            
    except Exception as e:
        print(f"Warning: Could not read page count for {ext}: {e}")
        return 1

def create_extraction_chain():
    """Extracts the core concepts from the document."""
    llm = ChatOpenAI(model="gpt-4.1-mini", temperature=0.3)
    
    system_instructions = """### ROLE
You are a Meta-Expert Analyst. Identify the most critical concepts from a document suitable for high-level testing.
### OUTPUT FORMAT
Return a simple JSON list of strings: {"concepts": ["concept1", "concept2"]}"""

    def build_extraction_msg(inputs):
        return [
            SystemMessage(content=system_instructions),
            HumanMessage(content=[
                {"type": "text", "text": "Extract the core concepts from this file."},
                {"type": "file", "file": {"file_id": inputs["file_id"]}}
            ])
        ]

    return build_extraction_msg | llm | JsonOutputParser()

def create_generation_chain(num_questions):
    """Generates the quiz using the dynamically injected question count."""
    llm = ChatOpenAI(model="gpt-4o-mini", temperature=0.5, model_kwargs={"response_format": {"type": "json_object"}})
    parser = JsonOutputParser()
    
    # Calculate exact distribution favoring Easy -> Medium -> Hard
    base = num_questions // 3
    remainder = num_questions % 3
    
    easy_qty = base + (1 if remainder >= 1 else 0)
    medium_qty = base + (1 if remainder == 2 else 0)
    hard_qty = base
    
    # The prompt now receives exact integers instead of fractions
    system_instructions = f"""You are a Senior Instructional Designer and Subject Matter Expert. Your goal is to create active recall assessment materials that help students master concepts from their study documents.

### OBJECTIVE
Analyze the provided user text/document and generate a multiple-choice quiz. The quiz must assess the user's understanding of the core concepts found strictly within the text.

### PEDAGOGICAL GUIDELINES
1. **Difficulty Distribution:** You must generate exactly {num_questions} questions in total, distributed EXACTLY as follows:
   - {easy_qty} Easy questions (direct recall of facts).
   - {medium_qty} Medium questions (application of concepts).
   - {hard_qty} Hard questions (analysis/evaluation based on Bloom's Taxonomy).
2. **Ordering:** You MUST present the questions strictly in ascending order of difficulty (Easy -> Medium -> Hard). Assign the correct "difficulty" label to each.
3. **Distractors:** The wrong options (distractors) must be plausible but clearly incorrect based on the text. Avoid obvious joke answers.
4. **Explanation Formatting:** The "explanation" field is critical. To maximize readability, you MUST separate the explanation of the correct answer and the breakdowns of each wrong option using double newlines (\\n\\n).

### STRICT CONSTRAINTS
1. **Source Truth:** Answer ONLY using the provided document. Do not use outside knowledge. If the information is not in the text, do not invent a question about it.
2. **Output Format:** You must output valid JSON only. Do not output conversational text before or after the JSON.

### FEW-SHOT EXAMPLE
{{
 "quiz_title": "Topic Summary",
 "questions": [
    {{
      "id": 1,
      "difficulty": "Easy",
      "question_text": "According to the document, why is active recall preferred over passive reading?",
      "options": [
        "A) It requires less mental effort.",
        "B) It strengthens neural pathways through testing.",
        "C) It allows students to read faster.",
        "D) It eliminates the need for textbooks."
      ],
      "correct_option": "B",
      "explanation": "The text states that self-testing (active recall) strengthens neural pathways (Roediger & Karpicke, 2006).\\n\\nOption A is incorrect because active recall explicitly requires more mental effort.\\n\\nOption C is incorrect because the text focuses on retention, not reading speed.\\n\\nOption D is incorrect as textbooks are still needed as source material."
    }}
  ]
}}

### FINAL INSTRUCTION
Generate the JSON format quiz now based on the attached document.
"""

    def build_generation_msg(inputs):
        file_id = inputs["file_id"]
        concepts = inputs["concepts_list"]
        
        return [
            SystemMessage(content=system_instructions),
            HumanMessage(content=[
                {"type": "text", "text": f"Focus the quiz strictly on these extracted core concepts: {concepts}"},
                {"type": "file", "file": {"file_id": file_id}}
            ])
        ]

    return build_generation_msg | llm | parser
